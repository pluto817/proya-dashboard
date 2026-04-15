import sys
from docx import Document

try:
    doc = Document(r'c:\Users\邓慧文\Desktop\珀莱雅可视化大屏\data\数据集\🔥 最终定稿・双碳深度融合版.docx')
    
    print("=== 文档标题 ===")
    print(doc.core_properties.title if doc.core_properties.title else "无标题")
    print()
    
    print("=== 文档内容 ===")
    for para in doc.paragraphs:
        if para.text.strip():
            print(para.text)
            print()
    
    print("\n=== 表格内容 ===")
    for i, table in enumerate(doc.tables):
        print(f"\n表格 {i+1}:")
        for row in table.rows:
            row_data = [cell.text.strip() for cell in row.cells]
            print("\t".join(row_data))
            
except Exception as e:
    print(f"错误: {e}")
    import traceback
    traceback.print_exc()
